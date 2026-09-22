import io
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_cover_letter_docx(cover_letter_text: str, candidate_name: str = "Applicant") -> io.BytesIO:
    """
    Creates a clean, professionally formatted DOCX file containing the cover letter.
    """
    doc = docx.Document()
    
    # Page Margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Style definitions
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1F, 0x24, 0x2E)

    # Candidate Header
    p_header = doc.add_paragraph()
    run_name = p_header.add_run(candidate_name)
    run_name.bold = True
    run_name.font.size = Pt(14)
    run_name.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    p_header.paragraph_format.space_after = Pt(18)

    # Body paragraphs
    paragraphs = [p.strip() for p in cover_letter_text.split("\n\n") if p.strip()]
    for p_content in paragraphs:
        p = doc.add_paragraph(p_content)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(12)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
