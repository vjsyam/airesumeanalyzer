import os
import docx
from sample_data import SAMPLE_RESUMES

def generate_samples():
    out_dir = os.path.join(os.path.dirname(__file__), "sample_resumes")
    os.makedirs(out_dir, exist_ok=True)
    
    # 1. Generate DOCX sample
    doc = docx.Document()
    morgan = SAMPLE_RESUMES["fullstack_engineer"]["text"]
    for para in morgan.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    docx_path = os.path.join(out_dir, "Morgan_Reed_Resume.docx")
    doc.save(docx_path)
    print(f"Generated DOCX: {docx_path}")

    # 2. Generate Alex Vance DOCX as well for easy file-drop testing
    doc2 = docx.Document()
    alex = SAMPLE_RESUMES["backend_engineer"]["text"]
    for para in alex.split("\n\n"):
        if para.strip():
            doc2.add_paragraph(para.strip())
    docx_path2 = os.path.join(out_dir, "Alex_Vance_Resume.docx")
    doc2.save(docx_path2)
    print(f"Generated DOCX: {docx_path2}")

if __name__ == "__main__":
    generate_samples()
